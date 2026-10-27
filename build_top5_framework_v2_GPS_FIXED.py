from pathlib import Path
from shutil import copy2
from collections import defaultdict
from statistics import median
import math

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


# ============================================================
# CONFIG
# ============================================================

BASE_DIR = Path(__file__).resolve().parent

INPUT = (BASE_DIR / "Data" / "Raw" / "Data_input.xlsx") if (BASE_DIR / "Data" / "Raw" / "Data_input.xlsx").exists() else BASE_DIR / "Data_input.xlsx"
BACKUP = INPUT.with_name("Data_input_BACKUP_BEFORE_TOP5_V3_GPS.xlsx")
OUTPUT = INPUT.with_name("Data_input_TOP5_FRAMEWORK_V3.xlsx")

SOURCE_SHEET = "09_District_Opportunity_Base"

FRAMEWORK_SHEET = "Top5_Scoring_Framework"
DEFINITION_SHEET = "Top5_Metric_Definition"


# Đúng 19 khu trong old-project universe
DISTRICTS_19 = [
    "Thủ Đức",
    "Tân Bình",
    "Bình Thạnh",
    "Quận 10",
    "Quận 1",
    "Quận 12",
    "Quận 7",
    "Bình Tân",
    "Gò Vấp",
    "Quận 3",
    "Tân Phú",
    "Dĩ An",
    "Quận 4",
    "Quận 6",
    "Thủ Dầu Một",
    "Vũng Tàu",
    "Bà Rịa",
    "Phú Nhuận",
    "Quận 5",
]


# ============================================================
# EXACT COLUMN MAPPING
# Theo đúng sheet 09 mà bạn vừa chụp
# ============================================================

COL = {
    "district_old": 1,              # A
    "competitor_locations": 3,      # C
    "competitor_brands": 4,         # D
    "icare_locations": 5,           # E
    "median_nearest_icare_km": 6,   # F
    "uncovered_share_5km": 7,       # G

    # H = population_reference_year => KHÔNG dùng làm population
    "population": 9,                # I official_population
    "density": 10,                  # J official_density

    "enterprise_count": 11,         # K
    "turnover": 12,                 # L
    "employee_compensation": 13,    # M

    "traffic_evidence": 14,         # N traffic_evidence_rows
}


# ============================================================
# HELPERS
# ============================================================

def to_num(value):
    """
    Convert Excel value to float.
    Missing/text flags remain None.
    """
    if value is None or value == "":
        return None

    if isinstance(value, bool):
        return int(value)

    if isinstance(value, (int, float)):
        if isinstance(value, float) and math.isnan(value):
            return None
        return float(value)

    s = str(value).strip()

    if s.upper() in {
        "",
        "NA",
        "N/A",
        "NO",
        "YES",
        "-",
        "NULL",
        "NONE",
    }:
        return None

    s = s.replace(",", "")

    if s.endswith("%"):
        s = s[:-1].strip()

    try:
        return float(s)
    except ValueError:
        raise ValueError(f"Không chuyển được thành số: {value!r}")


def normalize_share(value):
    """
    Source hiện đang lưu uncovered share dạng:
    0
    50
    100

    Output chuẩn:
    0
    0.5
    1.0
    """
    x = to_num(value)

    if x is None:
        return None

    if 0 <= x <= 1:
        return x

    if 1 < x <= 100:
        return x / 100

    raise ValueError(
        f"uncovered_share_5km bất thường: {x}. "
        "Chỉ chấp nhận range 0..1 hoặc 0..100."
    )


def style_header(ws):
    fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    font = Font(
        color="FFFFFF",
        bold=True,
    )

    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def autofit(ws, max_width=38):
    for col_idx, cells in enumerate(ws.columns, start=1):
        width = 10

        for cell in cells:
            value = "" if cell.value is None else str(cell.value)
            width = max(
                width,
                min(len(value) + 2, max_width),
            )

        ws.column_dimensions[
            get_column_letter(col_idx)
        ].width = width


# ============================================================
# CHECK INPUT
# ============================================================

if not INPUT.exists():
    raise FileNotFoundError(
        f"Không tìm thấy file:\n{INPUT}\n\n"
        "Hãy đặt script chung thư mục với Data_input.xlsx."
    )


print(f"[1] Input  : {INPUT.name}")


# ============================================================
# BACKUP
# ============================================================

if not BACKUP.exists():
    copy2(INPUT, BACKUP)
else:
    print("[2] Existing backup preserved (not overwritten)")

print(f"[2] Backup : {BACKUP.name}")


# ============================================================
# LOAD WORKBOOK
# ============================================================

wb = load_workbook(INPUT)

if SOURCE_SHEET not in wb.sheetnames:
    raise KeyError(
        f"Không có sheet '{SOURCE_SHEET}'."
    )

ws = wb[SOURCE_SHEET]

# ============================================================
# GPS SOURCE -> COVERAGE -> DISTRICT OPPORTUNITY
# Never read stale coverage metrics from the original workbook.
# ============================================================
def hdist_km(lat1, lon1, lat2, lon2):
    r = 6371.0088
    p1, p2 = math.radians(lat1), math.radians(lat2)
    dp = p2 - p1
    dl = math.radians(lon2 - lon1)
    a = math.sin(dp / 2)**2 + math.cos(p1)*math.cos(p2)*math.sin(dl / 2)**2
    return 2*r*math.asin(min(1, math.sqrt(a)))

market_sheet = wb['04_Market_Locations']
coverage_sheet = wb['06_iCare_Coverage_Evidence']
market_headers = {str(c.value).strip():c.column for c in market_sheet[1] if c.value is not None}
required = {'branch_id','brand','branch_name','address','district_old','latitude','longitude','analysis_role'}
if required - set(market_headers):
    raise ValueError(f"Market sheet thiếu cột: {required - set(market_headers)}")

items = []
for row in range(2, market_sheet.max_row + 1):
    obj = {h: market_sheet.cell(row, col).value for h,col in market_headers.items()}
    if not obj.get('branch_id'): continue
    obj['latitude'] = to_num(obj['latitude'])
    obj['longitude'] = to_num(obj['longitude'])
    if obj['latitude'] is None or obj['longitude'] is None:
        raise ValueError(f"Thiếu GPS cho {obj['branch_id']}")
    if not (-90 <= obj['latitude'] <= 90 and -180 <= obj['longitude'] <= 180):
        raise ValueError(f"GPS ngoài khoảng hợp lệ: {obj['branch_id']}")
    items.append(obj)

hubs = [x for x in items if x['analysis_role'] == 'ICARE_EXISTING']
competitors = [x for x in items if x['analysis_role'] == 'COMPETITOR_EVIDENCE']
unknown = [x for x in items if x['analysis_role'] not in ('ICARE_EXISTING','COMPETITOR_EVIDENCE')]
if unknown or len(hubs) != 3 or len(competitors) != 62:
    raise ValueError(f"Không đúng baseline footprint: hubs={len(hubs)}, competitors={len(competitors)}, unknown={len(unknown)}")
if len({x['branch_id'] for x in items}) != len(items):
    raise ValueError('branch_id trùng trong 04_Market_Locations')

# Replace coverage detail rows, preserve exact existing header names.
ch = [c.value for c in coverage_sheet[1]]
expected_ch = ['branch_id','brand','branch_name','district_old','address','latitude','longitude','nearest_icare_km','beyond_3km','beyond_5km','nearest_icare_branch','evidence_role','method_guardrail']
if ch != expected_ch:
    raise ValueError('06_iCare_Coverage_Evidence header thay đổi, cần audit lại mapping.')
if coverage_sheet.max_row > 1:
    coverage_sheet.delete_rows(2, coverage_sheet.max_row - 1)
bydistrict = defaultdict(list)
for x in competitors:
    d,h = min(((hdist_km(x['latitude'], x['longitude'], h['latitude'], h['longitude']),h) for h in hubs), key=lambda z:z[0])
    bydistrict[str(x['district_old']).strip()].append((d,x['brand']))
    coverage_sheet.append([x['branch_id'],x['brand'],x['branch_name'],x['district_old'],x['address'],x['latitude'],x['longitude'],d,'YES' if d>3 else 'NO','YES' if d>5 else 'NO',h['branch_name'],'MARKET_EVIDENCE_POINT','Haversine straight-line km; observed competitor point to nearest iCare; not district-wide demand or driving distance.'])

# Update only derived metrics, keep demographic/economic/traffic source columns unchanged.
headers09 = {str(c.value).strip(): c.column for c in ws[1] if c.value is not None}
needed09 = {'district_old','competitor_locations','competitor_brands','icare_locations','median_nearest_icare_km','uncovered_share_5km_pct'}
if needed09 - set(headers09):
    raise ValueError(f"09 sheet thiếu: {needed09 - set(headers09)}")
updated = {}
for row in range(2,ws.max_row + 1):
    district = ws.cell(row,headers09['district_old']).value
    if district not in DISTRICTS_19: continue
    points = bydistrict.get(district, [])
    expect_n = to_num(ws.cell(row,headers09['competitor_locations']).value)
    brands_n = to_num(ws.cell(row,headers09['competitor_brands']).value)
    actual_hubs = sum(x['district_old'] == district for x in hubs)
    old_hubs = to_num(ws.cell(row,headers09['icare_locations']).value)
    if expect_n != len(points) or brands_n != len({brand for d,brand in points}) or old_hubs != actual_hubs:
        raise ValueError(f"{district}: mismatch market footprint / brand / hubs; old=({expect_n},{brands_n},{old_hubs}) new=({len(points)},{len(set(b for d,b in points))},{actual_hubs})")
    med = median([d for d,brand in points]) if points else None
    beyond = 100*sum(d>5 for d,brand in points)/len(points) if points else None
    ws.cell(row,headers09['median_nearest_icare_km']).value = med
    ws.cell(row,headers09['uncovered_share_5km_pct']).value = beyond
    updated[district] = (med,beyond)
assert len(updated) == 19
print('[3] Coverage recalculated:',len(competitors),'competitors /',len(hubs),'hubs / 19 districts')
for district in ['Quận 7','Gò Vấp']:
    med,pct=updated[district]
    print(f'    {district}: median={med:.4f}km | outside5={pct:.2f}%')


# ============================================================
# FIND EXACT 19 DISTRICTS
# ============================================================

row_by_district = {}

for row in range(2, ws.max_row + 1):

    district = ws.cell(
        row=row,
        column=COL["district_old"],
    ).value

    if district is None:
        continue

    district = str(district).strip()

    # Chỉ lấy đúng 19 district
    # Không lấy summary rows kiểu HCMC 15 areas / ALL 19
    if district in DISTRICTS_19:

        if district in row_by_district:
            raise ValueError(
                f"District bị trùng trong source sheet: {district}"
            )

        row_by_district[district] = row


missing_districts = [
    d
    for d in DISTRICTS_19
    if d not in row_by_district
]

if missing_districts:
    raise ValueError(
        "Thiếu district trong source sheet:\n- "
        + "\n- ".join(missing_districts)
    )


# ============================================================
# BUILD RECORDS
# ============================================================

records = []


for district in DISTRICTS_19:

    row = row_by_district[district]

    competitor_locations = to_num(
        ws.cell(
            row=row,
            column=COL["competitor_locations"],
        ).value
    )

    competitor_brands = to_num(
        ws.cell(
            row=row,
            column=COL["competitor_brands"],
        ).value
    )

    icare_locations = to_num(
        ws.cell(
            row=row,
            column=COL["icare_locations"],
        ).value
    )

    median_nearest_icare_km = to_num(
        ws.cell(
            row=row,
            column=COL["median_nearest_icare_km"],
        ).value
    )

    uncovered_share_5km = normalize_share(
        ws.cell(
            row=row,
            column=COL["uncovered_share_5km"],
        ).value
    )

    population = to_num(
        ws.cell(
            row=row,
            column=COL["population"],
        ).value
    )

    density = to_num(
        ws.cell(
            row=row,
            column=COL["density"],
        ).value
    )

    traffic_evidence = to_num(
        ws.cell(
            row=row,
            column=COL["traffic_evidence"],
        ).value
    )

    enterprise_count = to_num(
        ws.cell(
            row=row,
            column=COL["enterprise_count"],
        ).value
    )

    turnover = to_num(
        ws.cell(
            row=row,
            column=COL["turnover"],
        ).value
    )

    employee_compensation = to_num(
        ws.cell(
            row=row,
            column=COL["employee_compensation"],
        ).value
    )


    # ========================================================
    # iCare presence
    # ========================================================

    if (
        icare_locations is not None
        and icare_locations > 0
    ):
        icare_presence = 1
    else:
        icare_presence = 0


    # ========================================================
    # SANITY CHECK
    # ========================================================

    # Tránh lỗi cũ: population lấy nhầm năm 2023
    if (
        population is not None
        and population < 50000
    ):
        raise ValueError(
            f"{district}: population={population:,.0f} "
            "quá nhỏ/bất thường.\n"
            "Có thể đang lấy nhầm population_reference_year."
        )


    if (
        density is not None
        and density <= 0
    ):
        raise ValueError(
            f"{district}: density={density} không hợp lệ."
        )


    if (
        competitor_locations is not None
        and competitor_locations < 0
    ):
        raise ValueError(
            f"{district}: competitor_locations âm."
        )


    if (
        competitor_brands is not None
        and competitor_brands < 0
    ):
        raise ValueError(
            f"{district}: competitor_brands âm."
        )


    if (
        median_nearest_icare_km is not None
        and median_nearest_icare_km < 0
    ):
        raise ValueError(
            f"{district}: median distance âm."
        )


    # ========================================================
    # COVERAGE N/A RULE
    # ========================================================

    # Quận 5:
    # competitor_locations = 0
    # nên median distance/share không tính được.
    # Đây là N/A hợp lệ, KHÔNG phải missing data.
    coverage_na = (
        competitor_locations == 0
        and median_nearest_icare_km is None
        and uncovered_share_5km is None
    )


    # ========================================================
    # CORE COMPLETENESS
    # ========================================================

    core_values = {
        "icare_presence": icare_presence,
        "competitor_locations": competitor_locations,
        "competitor_brands": competitor_brands,
        "population": population,
        "density": density,
        "traffic_evidence": traffic_evidence,
    }


    # Coverage chỉ bắt buộc nếu có observed competitor point
    if not coverage_na:

        core_values[
            "median_nearest_icare_km"
        ] = median_nearest_icare_km

        core_values[
            "uncovered_share_5km"
        ] = uncovered_share_5km


    core_available_n = sum(
        value is not None
        for value in core_values.values()
    )

    core_required_n = len(core_values)

    core_complete_pct = (
        core_available_n / core_required_n
        if core_required_n > 0
        else None
    )


    # ========================================================
    # ECONOMIC COMPLETENESS
    # ========================================================

    economic_values = [
        enterprise_count,
        turnover,
        employee_compensation,
    ]

    economic_complete_n = sum(
        value is not None
        for value in economic_values
    )


    # ========================================================
    # STATUS
    # ========================================================

    if core_complete_pct == 1:

        if economic_complete_n == 3:
            data_status = "CORE_READY_ECON_READY"

        else:
            data_status = "CORE_READY_ECON_PARTIAL"

    else:
        data_status = "NEED_RESEARCH"


    # ========================================================
    # NOTES
    # ========================================================

    notes = []


    if coverage_na:
        notes.append(
            "Coverage = N/A vì competitor_locations=0; "
            "không coi là missing và không đổi thành 0."
        )


    if economic_complete_n < 3:
        notes.append(
            "Economic context chưa đủ 3/3; "
            "chỉ dùng support/tie-break nếu chưa comparable."
        )


    notes.append(
        "traffic_evidence = evidence breadth, "
        "KHÔNG phải traffic volume/AADT."
    )


    # ========================================================
    # RECORD
    # ========================================================

    records.append(
        {
            "district_old": district,

            "icare_presence": icare_presence,

            "median_nearest_icare_km":
                median_nearest_icare_km,

            "uncovered_share_5km":
                uncovered_share_5km,

            "competitor_locations":
                competitor_locations,

            "competitor_brands":
                competitor_brands,

            "population":
                population,

            "density":
                density,

            "traffic_evidence":
                traffic_evidence,

            "enterprise_count":
                enterprise_count,

            "turnover":
                turnover,

            "employee_compensation":
                employee_compensation,

            "coverage_status":
                (
                    "N/A_NO_OBSERVED_COMPETITOR_POINTS"
                    if coverage_na
                    else "REQUIRED"
                ),

            "core_available_n":
                core_available_n,

            "core_required_n":
                core_required_n,

            "core_complete_pct":
                core_complete_pct,

            "economic_complete_n":
                economic_complete_n,

            "data_status":
                data_status,

            "source_trace":
                (
                    "GPS-derived 04 -> 06 -> 09_District_Opportunity_Base "
                    "| exact columns "
                    "A,C,D,E,F,G,I,J,K,L,M,N"
                ),

            "notes":
                " ".join(notes),
        }
    )


# ============================================================
# DELETE OLD GENERATED SHEETS
# ============================================================

for sheet_name in [
    FRAMEWORK_SHEET,
    DEFINITION_SHEET,
]:
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]


# ============================================================
# CREATE FRAMEWORK SHEET
# ============================================================

out = wb.create_sheet(
    FRAMEWORK_SHEET
)


headers = [
    "district_old",
    "icare_presence",
    "median_nearest_icare_km",
    "uncovered_share_5km",
    "competitor_locations",
    "competitor_brands",
    "population",
    "density",
    "traffic_evidence",
    "enterprise_count",
    "turnover",
    "employee_compensation",
    "coverage_status",
    "core_available_n",
    "core_required_n",
    "core_complete_pct",
    "economic_complete_n",
    "data_status",
    "source_trace",
    "notes",
]


out.append(headers)


for rec in records:

    out.append(
        [
            rec["district_old"],
            rec["icare_presence"],
            rec["median_nearest_icare_km"],
            rec["uncovered_share_5km"],
            rec["competitor_locations"],
            rec["competitor_brands"],
            rec["population"],
            rec["density"],
            rec["traffic_evidence"],
            rec["enterprise_count"],
            rec["turnover"],
            rec["employee_compensation"],
            rec["coverage_status"],
            rec["core_available_n"],
            rec["core_required_n"],
            rec["core_complete_pct"],
            rec["economic_complete_n"],
            rec["data_status"],
            rec["source_trace"],
            rec["notes"],
        ]
    )


style_header(out)


# ============================================================
# FORMATTING
# ============================================================

header_index = {
    cell.value: cell.column
    for cell in out[1]
}


for row in range(
    2,
    out.max_row + 1,
):

    out.cell(
        row,
        header_index[
            "median_nearest_icare_km"
        ],
    ).number_format = "0.00"


    out.cell(
        row,
        header_index[
            "uncovered_share_5km"
        ],
    ).number_format = "0.0%"


    out.cell(
        row,
        header_index["population"],
    ).number_format = "#,##0"


    out.cell(
        row,
        header_index["density"],
    ).number_format = "#,##0"


    out.cell(
        row,
        header_index["enterprise_count"],
    ).number_format = "#,##0"


    out.cell(
        row,
        header_index["turnover"],
    ).number_format = "#,##0"


    out.cell(
        row,
        header_index[
            "employee_compensation"
        ],
    ).number_format = "0.000"


    out.cell(
        row,
        header_index[
            "core_complete_pct"
        ],
    ).number_format = "0%"


autofit(out)


out.column_dimensions[
    get_column_letter(
        header_index["source_trace"]
    )
].width = 55


out.column_dimensions[
    get_column_letter(
        header_index["notes"]
    )
].width = 60


# ============================================================
# CREATE METRIC DEFINITION SHEET
# ============================================================

defs = wb.create_sheet(
    DEFINITION_SHEET
)


defs.append(
    [
        "metric",
        "group",
        "definition",
        "direction_for_selection",
        "use_in_scoring",
        "guardrail",
    ]
)


definition_rows = [

    [
        "icare_presence",
        "Hiện hữu iCare",
        "1 nếu khu đã có iCare hiện hữu, 0 nếu chưa có.",
        "0 thường thể hiện white-space lớn hơn.",
        "GUARDRAIL_ONLY",
        "Expansion eligibility only; not averaged into iCare score.",
    ],

    [
        "median_nearest_icare_km",
        "Hiện hữu iCare",
        "Khoảng cách trung vị từ observed competitor points tới iCare gần nhất.",
        "Cao hơn = observed coverage gap lớn hơn.",
        "CORE_WHEN_APPLICABLE",
        "Không phải district-wide coverage.",
    ],

    [
        "uncovered_share_5km",
        "Hiện hữu iCare",
        "Tỷ lệ observed competitor points nằm ngoài bán kính 5 km.",
        "Cao hơn = coverage gap lớn hơn.",
        "CORE_WHEN_APPLICABLE",
        "Observed points only; missing != 0.",
    ],

    [
        "competitor_locations",
        "Đối thủ / thị trường",
        "Số physical competitor locations đã verify.",
        "Cao hơn = market activity quan sát được mạnh hơn.",
        "CORE",
        "Không đồng nghĩa demand/customer count.",
    ],

    [
        "competitor_brands",
        "Đối thủ / thị trường",
        "Số brand đối thủ có physical presence đã verify.",
        "Cao hơn = breadth thị trường cạnh tranh lớn hơn.",
        "CORE",
        "Không phải market size trực tiếp.",
    ],

    [
        "population",
        "Business fit",
        "Dân số chính thức của khu.",
        "Cao hơn = potential catchment lớn hơn.",
        "CORE",
        "Potential base, không phải demand thực tế.",
    ],

    [
        "density",
        "Business fit",
        "Mật độ dân số của khu.",
        "Cao hơn = catchment dày hơn.",
        "CORE",
        "Không phản ánh willingness to pay.",
    ],

    [
        "traffic_evidence",
        "Business fit",
        "Số/breadth bằng chứng traffic đã thu thập.",
        "Cao hơn = evidence breadth mạnh hơn.",
        "GUARDRAIL_ONLY",
        "Không gọi là traffic volume hoặc AADT.",
    ],

    [
        "enterprise_count",
        "Economic context",
        "Số doanh nghiệp đang hoạt động.",
        "Cao hơn có thể phản ánh business activity.",
        "SUPPORT",
        "Chỉ đưa vào score nếu comparable toàn bộ universe.",
    ],

    [
        "turnover",
        "Economic context",
        "Doanh thu doanh nghiệp.",
        "Cao hơn có thể phản ánh economic activity.",
        "SUPPORT",
        "Không mix district-level với province-level.",
    ],

    [
        "employee_compensation",
        "Economic context",
        "Thu nhập/compensation bình quân lao động doanh nghiệp.",
        "Cao hơn có thể phản ánh purchasing-power/business context.",
        "SUPPORT",
        "Không phải household income.",
    ],
]


for row in definition_rows:
    defs.append(row)


style_header(defs)

autofit(
    defs,
    max_width=50,
)


defs.column_dimensions["C"].width = 60
defs.column_dimensions["F"].width = 60


# ============================================================
# FINAL VALIDATION
# ============================================================

# 19 district + 1 header
if out.max_row != 20:
    raise AssertionError(
        "Framework phải có đúng "
        "19 district + 1 header = 20 dòng.\n"
        f"Hiện tại có: {out.max_row}"
    )


districts_written = [
    out.cell(
        row=row,
        column=1,
    ).value
    for row in range(
        2,
        21,
    )
]


if districts_written != DISTRICTS_19:
    raise AssertionError(
        "Danh sách 19 district không đúng."
    )


# Chống summary rows
bad_rows = [
    "HCMC 15 areas",
    "ALL 19",
    "Quận 1 / Quận 3 / Tân Bình / Phú Nhuận",
]


for bad in bad_rows:

    if bad in districts_written:
        raise AssertionError(
            f"Summary row bị kéo nhầm: {bad}"
        )


# ============================================================
# SAVE
# ============================================================

wb.save(OUTPUT)


# ============================================================
# PRINT RESULT
# ============================================================

print()
print("=" * 70)
print("DONE - TOP5 FRAMEWORK V3")
print("=" * 70)

print(
    f"Output : {OUTPUT.name}"
)

print(
    f"Backup : {BACKUP.name}"
)

print(
    f"Rows   : {len(records)} districts"
)

print()
print("DATA COMPLETENESS:")
print()


for rec in records:

    print(
        f"- {rec['district_old']}: "
        f"core "
        f"{rec['core_available_n']}/"
        f"{rec['core_required_n']} "
        f"| econ "
        f"{rec['economic_complete_n']}/3 "
        f"| {rec['data_status']}"
    )


print()
print("=" * 70)

print(
    "CHECK QUAN TRỌNG:"
)

print(
    "- population lấy từ cột I = official_population."
)

print(
    "- KHÔNG lấy cột H = population_reference_year."
)

print(
    "- uncovered_share_5km: 50 -> 50%, 100 -> 100%."
)

print(
    "- chỉ giữ đúng 19 district."
)

print(
    "- summary rows bị loại."
)

print(
    "- Quận 5 coverage có thể N/A hợp lệ nếu competitor_locations=0."
)

print(
    "- economic data thiếu không bị đổi thành 0."
)

print(
    "- traffic_evidence chỉ là evidence breadth, không phải traffic volume."
)

print("=" * 70)
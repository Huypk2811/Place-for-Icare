from pathlib import Path
from shutil import copy2
from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent

INPUT = BASE_DIR / "Data_input.xlsx"
BACKUP = BASE_DIR / "Data_input_BACKUP_BEFORE_DENSITY_UPDATE.xlsx"

SHEET_NAME = "09_District_Opportunity_Base"

# Cột theo sheet hiện tại:
# A = district_old
# J = official_density
DISTRICT_COL = 1
DENSITY_COL = 10

UPDATES = {
    "Vũng Tàu": 2422,
    "Bà Rịa": 1188,
}

if not INPUT.exists():
    raise FileNotFoundError(
        f"Không tìm thấy {INPUT.name} trong thư mục {BASE_DIR}"
    )

# Backup trước
copy2(INPUT, BACKUP)

wb = load_workbook(INPUT)

if SHEET_NAME not in wb.sheetnames:
    raise KeyError(f"Không tìm thấy sheet: {SHEET_NAME}")

ws = wb[SHEET_NAME]

found = set()

for row in range(2, ws.max_row + 1):
    district = ws.cell(row=row, column=DISTRICT_COL).value

    if district is None:
        continue

    district = str(district).strip()

    if district in UPDATES:
        old_value = ws.cell(row=row, column=DENSITY_COL).value
        new_value = UPDATES[district]

        ws.cell(row=row, column=DENSITY_COL).value = new_value

        print(
            f"{district}: official_density "
            f"{old_value} -> {new_value}"
        )

        found.add(district)

missing = set(UPDATES.keys()) - found

if missing:
    raise ValueError(
        "Không tìm thấy district: "
        + ", ".join(sorted(missing))
    )

wb.save(INPUT)

print()
print("DONE")
print(f"Updated: {INPUT.name}")
print(f"Backup : {BACKUP.name}")